import csv
import hashlib
import json
import math
import re
import statistics
import unicodedata
from datetime import date
from pathlib import Path
from openpyxl import load_workbook

ALIASES = {
 'keyword':['关键词','流量词','搜索词','Keyword','Search Term','Keyword Phrase'],
 'translation':['关键词翻译','翻译','Translation'],
 'search_volume':['月搜索量','Search Volume','Monthly Search Volume'],
 'aba_rank':['ABA周排名','ABA周排行','ABA Rank','ABA Weekly Rank'],
 'purchase_rate':['购买率','Purchase Rate'],
 'ppc':['PPC价格','PPC竞价','PPC Bid','PPC'],
 'traffic_share':['流量占比','Traffic Share'],
 'organic_rank':['自然排名','Organic Rank'],
 'ad_rank':['广告排名','Sponsored Rank'],
 'ad_competitors':['近7天广告竞品数','广告竞品数','Sponsored ASINs'],
 'related_asins':['相关ASIN','Related ASINs'],
 'related_count':['相关产品','Related Products'],
 'asin':['ASIN','Competitor ASIN'],
}
METRICS=['search_volume','aba_rank','purchase_rate','ppc','traffic_share','organic_rank','ad_rank','ad_competitors','related_count']
GLOBAL=['search_volume','aba_rank','purchase_rate','ppc','ad_competitors']
CURRENCY={'US':'USD','DE':'EUR','UK':'GBP','CA':'CAD','FR':'EUR','IT':'EUR','ES':'EUR','JP':'JPY','AU':'AUD'}
ASIN = re.compile(r'\b[A-Z0-9]{10}\b')

def norm(value): return ' '.join(unicodedata.normalize('NFKC',str(value)).casefold().split())
def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def dump(path,data): Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def load(path): return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def number(value,key,market):
    if value is None or str(value).strip() in {'','-','—','N/A','n/a'}: return None,'blank'
    if key in {'organic_rank','ad_rank'} and '无排名' in str(value): return None,'not_ranked'
    if isinstance(value,bool): raise ValueError('布尔值不是数值')
    text=str(value).strip().replace(',','')
    if key=='ppc':
        symbols={'€':'EUR','£':'GBP','¥':'JPY','$':CURRENCY[market] if market in {'US','CA','AU'} else 'USD'}
        for sym,cur in symbols.items():
            if sym in text:
                if cur!=CURRENCY[market]: raise ValueError('币种与站点不匹配')
                text=text.replace(sym,'')
    percent=text.endswith('%')
    if percent: text=text[:-1]
    value=float(text)
    if percent: value/=100
    if not math.isfinite(value) or value<0: raise ValueError('非有限值或负值')
    if key in {'purchase_rate','traffic_share'} and value>1: raise ValueError('比例必须明确为0–1或百分数')
    return value,'observed'

def read_sources(input_dir,market,manifest=None,mapping=None):
    root=Path(input_dir)
    if root.is_dir():
        unsupported=[p.name for p in root.iterdir() if p.is_file() and p.suffix.lower() in {'.xls','.xlsm','.xlsb','.ods','.zip'}]
        if unsupported: raise ValueError('输入含未支持的表格格式，请明确转换并核对字段: '+', '.join(unsupported))
    elif root.suffix.lower() not in {'.xlsx','.csv','.tsv'}:
        raise ValueError('未支持的输入格式: '+root.suffix)
    files=sorted(p for p in root.iterdir() if p.suffix.lower() in {'.xlsx','.csv','.tsv'} and not p.name.startswith('~$')) if root.is_dir() else [root]
    if not files: raise ValueError('没有输入表格')
    if market not in CURRENCY: raise ValueError('未支持的站点')
    aliases={norm(a):k for k,aa in ALIASES.items() for a in aa}
    for alias,key in (mapping or {}).items():
        if key not in ALIASES: raise ValueError('字段映射目标未知: '+key)
        aliases[norm(alias)]=key
    m=manifest
    manifest_sources={}
    if m:
        for key in ('batch_id','market','window_start','window_end','provider','definition'):
            if not m.get(key): raise ValueError('批次清单缺少 '+key)
        if m['market']!=market: raise ValueError('批次站点不匹配')
        if date.fromisoformat(m['window_start'])>date.fromisoformat(m['window_end']): raise ValueError('批次日期倒置')
        manifest_sources={s['file']:s for s in m['sources']}
        if len(manifest_sources)!=len(m['sources']) or set(manifest_sources)!={p.name for p in files}: raise ValueError('清单与输入文件不一致')
    rows=[]; diagnostic=[]; errors=[]
    for path in files:
        file_hash=digest(path)
        if m and manifest_sources[path.name].get('sha256')!=file_hash: raise ValueError('清单文件哈希不匹配: '+path.name)
        wb=None
        if path.suffix.lower()=='.xlsx':
            wb=load_workbook(path,read_only=True,data_only=True)
            sheets=[(ws.title,list(ws.iter_rows(values_only=True))) for ws in wb]
            wb.close()
        else:
            with path.open(encoding='utf-8-sig',newline='') as f: sheets=[('CSV',list(csv.reader(f,delimiter='\t' if path.suffix.lower()=='.tsv' else ',')))]
        roster=set()
        for name,values in sheets:
            if norm(name)=='asin':
                roster.update(a for row in values[1:] for cell in row if cell for a in ASIN.findall(str(cell).upper()))
        if m:
            declared=set(manifest_sources[path.name].get('asins',[]))
            if not declared or any(not ASIN.fullmatch(a) for a in declared): raise ValueError('清单 ASIN 无效')
            if roster and declared!=roster: raise ValueError('清单与表内 ASIN 清单不一致')
            roster=declared
        valid_sheets=0
        for name,values in sheets:
            header=None
            for index,line in enumerate(values[:30]):
                found={}
                for col,v in enumerate(line):
                    key=aliases.get(norm(v)) if v is not None else None
                    if key:
                        if key in found: errors.append(f'{path.name}/{name}:重复字段 {key}')
                        found[key]=col
                if 'keyword' in found:
                    header=(index,found,line); break
            if header is None:
                # Recognized auxiliary sheets are audited; unknown data sheets fail explicitly.
                known=norm(name) in {'unique words','notes','asin','说明','词频'}
                if not known and any(any(v is not None and v!='' for v in row) for row in values): errors.append(f'{path.name}/{name}:无法识别关键词列，请提供映射')
                diagnostic.append({'file':path.name,'sheet':name,'status':'auxiliary' if known else 'unrecognized'})
                continue
            valid_sheets+=1
            start,fields,headers=header
            sheet_asins=set(ASIN.findall(name.upper()))
            # The ASIN in a multi-ASIN title is only a display anchor.
            single=next(iter(sheet_asins)) if len(sheet_asins)==1 and not roster else None
            sheet_market=re.match(r'^(US|DE|UK|CA|FR|IT|ES|JP|AU)[-_]',name.upper())
            if sheet_market and sheet_market[1]!=market: errors.append(f'{path.name}/{name}:站点不匹配')
            for row_num,line in enumerate(values[start+1:],start+2):
                if not any(v is not None and v!='' for v in line): continue
                def cell(key): return line[fields[key]] if key in fields and fields[key]<len(line) else None
                keyword=cell('keyword')
                if keyword is None or not str(keyword).strip(): errors.append(f'{path.name}/{name}/{row_num}:数据行缺少关键词'); continue
                if not isinstance(keyword,str): errors.append(f'{path.name}/{name}/{row_num}:关键词非文本'); continue
                related=set(ASIN.findall(str(cell('related_asins') or '').upper()))
                explicit=set(ASIN.findall(str(cell('asin') or '').upper()))
                if explicit and roster and not explicit<=roster: errors.append(f'{path.name}/{name}/{row_num}:行ASIN不在批次清单')
                identity=explicit or ((related & roster) if roster else ({single} if single else related))
                if m and not identity: errors.append(f'{path.name}/{name}/{row_num}:无法识别目标ASIN')
                metrics={}; states={}
                for key in METRICS:
                    if key not in fields: metrics[key]=None; states[key]='column_missing'; continue
                    try: metrics[key],states[key]=number(cell(key),key,market)
                    except (ValueError,TypeError) as e:
                        metrics[key]=None; states[key]='invalid'; errors.append(f'{path.name}/{name}/{row_num}/{key}: {e} ({cell(key)!r})')
                raw=[{'column':i+1,'header':str(h),'value':line[i] if i<len(line) else None} for i,h in enumerate(headers) if h is not None]
                rid=hashlib.sha256(f'{file_hash}|{name}|{row_num}'.encode()).hexdigest()[:24]
                rows.append({'row_id':rid,'file':path.name,'file_sha256':file_hash,'sheet':name,'row':row_num,
                             'raw_keyword':keyword,'keyword':norm(keyword),'translation':cell('translation'),
                             'asins':sorted(identity),'related_asins':sorted(related),'target_roster':sorted(roster or ({single} if single else set())),
                             'metrics':metrics,'states':states,'raw':raw,'market':market,
                             'window':[m['window_start'],m['window_end']] if m else None})
            diagnostic.append({'file':path.name,'sheet':name,'status':'accepted','header_row':start+1,'fields':list(fields)})
        if not valid_sheets: errors.append(path.name+':没有可用关键词Sheet')
    if errors: raise ValueError('输入未通过检查:\n'+'\n'.join(errors[:40])+f'\n共{len(errors)}项；无数据被静默丢弃')
    if not rows: raise ValueError('没有关键词数据行')
    return rows,diagnostic

def aggregate(rows,verified):
    groups={}
    for r in rows: groups.setdefault(r['keyword'],[]).append(r)
    result=[]
    for word,rr in groups.items():
        representative=max(rr,key=lambda r:sum(r['metrics'][k] is not None for k in GLOBAL))
        metrics={k:representative['metrics'][k] for k in GLOBAL}
        if verified:
            for key in GLOBAL:
                values=[r['metrics'][key] for r in rr if r['metrics'][key] is not None]
                metrics[key]=statistics.median(values) if values else None
        asins=sorted({a for r in rr for a in r['asins']})
        source_counts=len({r['file'] for r in rr})
        metrics['verified_competitors']=len(asins) if verified else None
        metrics['observed_asins']=len(asins) if asins else None
        metrics['related_count']=representative['metrics']['related_count']
        mismatches=[k for k in GLOBAL if len({r['metrics'][k] for r in rr if r['metrics'][k] is not None})>1]
        result.append({'id':hashlib.sha256(word.encode()).hexdigest()[:20],'keyword':word,'translation':representative['translation'],
                       'row_ids':[r['row_id'] for r in rr],'asins':asins,'source_files':source_counts,'metrics':metrics,
                       'metric_row_id':None if verified else representative['row_id'],'metric_method':'median_same_batch' if verified else 'one_source_row',
                       'metric_differences':mismatches})
    if sum(len(r['row_ids']) for r in result)!=len(rows) or len({r['row_id'] for r in rows})!=len(rows): raise ValueError('源行血缘不守恒')
    return sorted(result,key=lambda r:r['keyword'])

