"""Portable Excel fallback only when @oai/artifact-tool is unavailable."""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

def clean(v): return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
def export(input_path,output_dir):
    payload=json.loads(Path(input_path).read_text(encoding='utf-8-sig'))
    if not payload.get('rows') or any(len(r)!=len(payload['headers']) for r in payload['rows']): raise ValueError('Excel payload无效')
    out=Path(output_dir);out.mkdir(parents=True,exist_ok=True)
    wb=Workbook();ws=wb.active;ws.title='关键词库';ws.sheet_view.showGridLines=False
    ws.append(payload['headers'])
    for row in payload['rows']:ws.append([clean(x) for x in row])
    ws.freeze_panes='A2'
    headers=PatternFill('solid',fgColor='233E47')
    for c in ws[1]:c.fill=headers;c.font=Font(color='FFFFFF',bold=True,name='Microsoft YaHei',size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=42
    widths={'priority':10,'keyword':35,'translation':22,'level':10,'tags':23,'search_volume':15,'aba_rank':15,'verified_competitors':19,'observed_asins':24,'related_count':20,'purchase_rate':16,'ppc':18,'ad_competitors':16,'match':19,'reason':60,'evidence':50}
    for col,key in enumerate(payload['keys'],1):
        ws.column_dimensions[get_column_letter(col)].width=widths.get(key,18)
        for cells in ws.iter_cols(min_col=col,max_col=col,min_row=2):
            for c in cells:
                c.font=Font(name='Microsoft YaHei',size=10)
                c.alignment=Alignment(wrap_text=key in {'keyword','translation','tags','reason','evidence','match'},vertical='center')
                if key=='purchase_rate':c.number_format='0.00%'
                elif key=='ppc':c.number_format='0.00'
                elif key in {'search_volume','aba_rank','verified_competitors','observed_asins','related_count','ad_competitors'}:c.number_format='#,##0'
    for row in range(2,ws.max_row+1):
        ws.row_dimensions[row].height=56
        priority=ws.cell(row,1);priority.fill=PatternFill('solid',fgColor={'P1':'DDEBDD','P2':'F7EDD5','P3':'EAE4F0'}.get(priority.value,'FFFFFF'));priority.font=Font(bold=True)
    table=Table(displayName='Keywords',ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=False,showColumnStripes=False)
    ws.add_table(table)
    if payload.get('notes'):
        note=wb.create_sheet('说明');note.column_dimensions['A'].width=110;note.append(['本次数据说明'])
        for n in payload['notes']:note.append([clean(n)])
        for r in note:
            r[0].alignment=Alignment(wrap_text=True,vertical='center');note.row_dimensions[r[0].row].height=55
        note['A1'].fill=headers;note['A1'].font=Font(color='FFFFFF',bold=True)
    dest=out/'关键词库.xlsx';wb.save(dest);wb.close()
    check=load_workbook(dest,read_only=True,data_only=False)
    if check['关键词库'].max_row!=len(payload['rows'])+1:raise ValueError('保存后行数不一致')
    check.close();print(dest)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('payload');p.add_argument('output_dir');a=p.parse_args();export(a.payload,a.output_dir)

